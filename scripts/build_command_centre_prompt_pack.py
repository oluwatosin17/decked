from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = "Decked_Command_Centre_Implementation_Prompt_Pack.docx"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(28, 28, 30)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12

for name, size, before, after in [
    ("Title", 30, 0, 14), ("Heading 1", 20, 18, 8),
    ("Heading 2", 15, 14, 6), ("Heading 3", 12, 10, 4),
]:
    st = styles[name]
    st.font.name = "Aptos Display"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

title_style_ppr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

if "Prompt Text" not in styles:
    ps = styles.add_style("Prompt Text", WD_STYLE_TYPE.PARAGRAPH)
else:
    ps = styles["Prompt Text"]
ps.font.name = "Aptos Mono"
ps.font.size = Pt(9)
ps.font.color.rgb = RGBColor(35, 35, 38)
ps.paragraph_format.left_indent = Inches(0.18)
ps.paragraph_format.right_indent = Inches(0.08)
ps.paragraph_format.space_after = Pt(7)
ps.paragraph_format.line_spacing = 1.08

def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def add_header_footer(sec):
    hp = sec.header.paragraphs[0]
    hp.text = "DECKED  |  COMMAND CENTRE BUILD PACK"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.name = "Aptos"
    hp.runs[0].font.size = Pt(8)
    hp.runs[0].font.bold = True
    hp.runs[0].font.color.rgb = RGBColor(90, 90, 94)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Decked internal implementation guide  •  ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)

add_header_footer(section)

title = doc.add_paragraph(style="Title")
title_run = title.add_run("Decked Command Centre Implementation Prompt Pack")
title_run.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = subtitle.add_run("A sequenced set of coding prompts for analytics instrumentation, the internal dashboard, testing, and launch")
r.font.name = "Aptos"
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(65, 65, 70)

doc.add_paragraph()
meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
meta.columns[0].width = Inches(1.55)
meta.columns[1].width = Inches(5.1)
for row in meta.rows:
    for cell in row.cells:
        set_cell_margins(cell, 100, 140, 100, 140)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
for i, (k, v) in enumerate([
    ("Product", "Decked"),
    ("Deliverable", "Internal command centre and analytics foundation"),
    ("Recommended use", "Run one numbered prompt at a time in the Decked repository"),
    ("Prepared", "9 October 2026"),
]):
    meta.cell(i,0).text = k
    meta.cell(i,1).text = v
    shade_cell(meta.cell(i,0), "242426")
    for run in meta.cell(i,0).paragraphs[0].runs:
        run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    if i % 2 == 1: shade_cell(meta.cell(i,1), "F4F4F5")

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Main recommendation  ")
r.bold = True
p.add_run("Build the measurement foundation before building charts. The dashboard will only be useful if event definitions, identity rules, and server-side multiplayer facts are reliable.")

doc.add_page_break()

doc.add_heading("How to use this pack", level=1)
doc.add_paragraph("These prompts are written for an AI coding agent with access to the Decked repository. Run them in order. Let the agent inspect the current code before it changes anything, and do not combine several implementation prompts into one turn. Smaller changes are easier to review, test, and reverse.")
for text in [
    "Start each new coding chat with the reusable context prompt, then provide the relevant numbered prompt.",
    "Require the agent to preserve unrelated work and report every file it changes.",
    "Review migrations before applying them to production. Database changes must be tested in a non-production Supabase project first.",
    "Commit or checkpoint the repository after each completed phase.",
    "If the agent discovers a conflict with the current architecture, ask it to explain the conflict and revise the plan before proceeding.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Suggested sequence", level=2)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [0.65, 1.55, 3.65, 1.0]
for i,w in enumerate(widths): table.columns[i].width = Inches(w)
hdr = table.rows[0].cells
for i, text in enumerate(["Phase", "Focus", "Prompts", "Gate"]):
    hdr[i].text = text; shade_cell(hdr[i], "242426"); set_cell_margins(hdr[i])
    for run in hdr[i].paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
rows = [
    ("1", "Discovery", "1 to 3", "Approved plan"),
    ("2", "Data foundation", "4 to 7", "Trusted events"),
    ("3", "Dashboard", "8 to 13", "Usable MVP"),
    ("4", "Quality and launch", "14 to 18", "Production ready"),
]
for ri, vals in enumerate(rows):
    cells=table.add_row().cells
    for i,v in enumerate(vals):
        cells[i].text=v; set_cell_margins(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ri%2: shade_cell(cells[i], "F5F7FA")

doc.add_heading("Before every prompt", level=2)
doc.add_paragraph("Replace bracketed placeholders such as [SUPABASE PROJECT] or [ADMIN EMAILS]. If you do not know a value, tell the coding agent to leave it as an environment variable and document it. Never paste production service-role keys into chat or commit them to the repository.")

doc.add_page_break()

context_prompt = """You are working in the existing Decked repository. Decked is a React 19 and TypeScript application built with Vite and deployed on Vercel. It uses Supabase authentication, Postgres, RPC functions, and Realtime for separate-device multiplayer. It also supports pass-and-play on one device. The current application exposes 21 multiplayer game IDs and uses anonymous Supabase users for players.

Before making changes, inspect the repository, relevant migrations, package scripts, and any repository instructions. Preserve unrelated user changes. Use the existing design language and conventions. Do not invent files, tables, environment variables, or APIs without checking first.

For every implementation task:
1. State the current behavior and the change you intend to make.
2. Identify security, privacy, migration, and backward-compatibility risks.
3. Implement the smallest complete change.
4. Add or update tests where the repository supports them.
5. Run the most relevant lint, type-check, build, and test commands.
6. Report changed files, verification results, assumptions, and remaining work.

Never expose Supabase service-role credentials to the browser. Never collect player display names, room codes, free-form answers, or prompt text in general analytics. Use stable identifiers. Keep development and production analytics separable."""

doc.add_heading("Reusable context prompt", level=1)
doc.add_paragraph("Paste this at the beginning of a new coding chat. You can omit it only when continuing in the same chat and the agent still has the repository context.")
doc.add_paragraph(context_prompt, style="Prompt Text")

prompts = [
("1", "Audit the current product and produce an implementation map", "Discovery", "A repository-grounded map of the current application and the safest implementation path.", """Audit the Decked repository for the Command Centre project. Do not modify code yet.

Inspect navigation, game selection, pass-and-play state, Quick Play, the multiplayer client, Supabase authentication, every migration, deployment configuration, and any existing analytics or error reporting. Confirm the canonical list of game IDs, category IDs, deck or mode IDs, room lifecycle transitions, and locations where a game can start, advance, finish, replay, or be abandoned.

Produce:
- a concise architecture map;
- a list of all instrumentation points with exact file and function references;
- an inventory of existing database tables and RPC functions relevant to analytics;
- risks caused by room deletion or expiry;
- gaps in tests and tooling;
- a phased implementation plan with migration order;
- decisions that require product input.

Do not propose a rewrite. Optimize for incremental changes to the current React, Vite, Supabase, and Vercel architecture. End with a recommended definition of MVP and an explicit list of files likely to change."""),
("2", "Define the measurement plan and metric contract", "Discovery", "One source of truth for events, properties, identities, and metrics.", """Using the completed Decked repository audit, write a measurement plan before implementing analytics.

Define the canonical event taxonomy for navigation, discovery, Quick Play, game setup, pass-and-play, multiplayer rooms, gameplay, completion, replay, rematch, host handoff, and errors. For each event specify: event name, trigger, required properties, optional properties, source of truth, deduplication approach, and whether it should be emitted by the browser, database, or both.

Define exact formulas for daily, weekly, and monthly active players; game starts; first-card rate; completion rate; median session duration; replay rate; multiplayer lobby conversion; room join success; multiplayer completion; and error-free sessions. Explain identity limitations created by anonymous authentication and browser storage.

Apply these privacy constraints: never collect display names, room codes, free-form answers, or prompt text; use stable game, deck, category, prompt, session, and app-version identifiers; use coarse country only if available server-side.

Deliver the measurement plan as a versioned Markdown specification in the repository. Do not implement tracking yet. Run any existing documentation checks and report unresolved decisions."""),
("3", "Create the technical design and implementation backlog", "Discovery", "An approved system design that prevents premature dashboard work.", """Create a technical design for the Decked Command Centre based on the repository audit and measurement plan. Do not implement production code yet.

Compare two viable analytics architectures: a Supabase-native append-only event pipeline and a managed analytics provider such as PostHog. Evaluate cost, operational burden, event validation, funnels, retention, data ownership, privacy, and integration effort for this codebase. Recommend one architecture for the MVP and explain why.

Design the selected system, including event ingestion, validation, idempotency, sessionization, authoritative server-side multiplayer facts, aggregate tables or materialized views, staff authentication, role authorization, audit logs, dashboard queries, data retention, environment separation, and failure handling.

Produce a dependency-ordered backlog with small implementation tickets. Each ticket must include acceptance criteria, affected areas, migration risk, verification steps, and rollback notes. Identify which tickets can be developed independently and where a human approval gate is required."""),
("4", "Add the analytics database foundation", "Data foundation", "Append-only storage with RLS, indexes, retention fields, and safe access paths.", """Implement the approved Supabase analytics database foundation.

Create a new timestamped migration containing the minimum normalized schema required for append-only analytics events, game sessions, multiplayer room facts, daily aggregates, staff roles, and staff audit logs. Use appropriate UUIDs, timestamps, constrained enums or checks, JSONB only where flexibility is intentional, indexes for expected dashboard filters, and idempotency keys for duplicate prevention.

Enable row-level security on every new table. Browser clients must not be able to read raw analytics or assign staff roles. Create narrowly scoped RPC or server-only paths for event ingestion. Keep service-role operations outside the browser. Do not alter or weaken the existing gameplay RLS policies.

Add comments explaining sensitive columns and retention expectations. If local Supabase testing is available, apply the migration and test expected grants and denials. Otherwise perform static SQL validation and provide exact non-production deployment steps. Update the technical documentation and report rollback instructions."""),
("5", "Implement the typed analytics client", "Data foundation", "A single safe client API used consistently across the application.", """Implement a typed analytics module for Decked using the approved measurement plan and database ingestion path.

Requirements:
- expose one small track function with TypeScript discrimination between event names and their permitted properties;
- generate and persist analytics_session_id and game_session_id values appropriately;
- attach app version, environment, device class, game ID, play mode, deck or category ID, and timestamps only when valid;
- batch or retry safely without blocking gameplay;
- deduplicate events using a generated event ID;
- fail silently for players while recording a development diagnostic;
- respect an environment flag that disables analytics;
- keep production and development data separated;
- never send display names, room codes, answers, or prompt text.

Add unit tests for schema validation, forbidden properties, identity persistence, disabled tracking, retry behavior, and event deduplication. Add concise developer documentation with examples. Run lint, type checking, tests, and a production build."""),
("6", "Instrument discovery and pass and play", "Data foundation", "End-to-end visibility from application open to local game completion.", """Instrument Decked's discovery and pass-and-play journeys using the typed analytics client. Follow the measurement plan exactly and avoid duplicate events caused by React rendering.

Cover application open, home and browse views, category filters, Quick Play views and shuffles, game selection, play-mode selection, deck or mode selection, local game configuration, game start, first card, card advance, completion, replay, and abandonment where it can be measured reliably.

Use shared components or hooks where they can normalize behavior without rewriting every game. Where individual games require adapters, create an explicit registry and document coverage. Do not send prompt text or player-entered content.

Add tests for the main journey and for duplicate prevention under React Strict Mode. Provide a coverage matrix for all 21 games showing which lifecycle events are implemented and any game-specific exceptions. Run lint, type checking, tests, and a production build."""),
("7", "Instrument authoritative multiplayer lifecycle events", "Data foundation", "Reliable room analytics that survive room deletion and client disconnects.", """Add authoritative multiplayer analytics without relying solely on browser events.

Update Supabase migrations and RPC functions so successful room creation, join, game start, card progression where needed, completion, player exit, room end, rematch, and host handoff write durable analytics facts transactionally. Record failure events at the safest available layer without exposing secrets or sensitive inputs. Ensure repeated RPC calls do not create duplicate facts.

Use restricted room references; do not store room codes or display names. Preserve the existing gameplay response shapes and authorization rules. Ensure analytics failure cannot break gameplay transactions unless data integrity would otherwise become misleading, and document that decision.

Add browser-side attempted and failed events only where the server cannot observe them, including invalid join attempts and Realtime subscription state. Test successful flows, authorization denials, retries, duplicate requests, expired rooms, host disconnect, and deletion of operational room data. Demonstrate that historical analytics remains available after a room ends."""),
("8", "Scaffold the protected Command Centre", "Dashboard", "A secure internal application shell and routing structure.", """Implement the protected Command Centre shell at /command-centre or the route selected in the technical design.

Create staff sign-in, sign-out, session restoration, loading, unauthorized, and error states. Authorization must be enforced by Supabase policies or a trusted server path, not only by frontend route guards. Anonymous gameplay accounts must not gain dashboard access.

Build a responsive desktop-first layout with navigation for Overview, Games, Funnels, Multiplayer, Reliability, and Settings. Match Decked's visual identity while prioritizing legibility: dark neutral surfaces, restrained red accent, readable tables, accessible focus states, and WCAG AA contrast.

Do not build charts yet. Create reusable page, filter, KPI, table, loading, empty, and error primitives. Add route and authorization tests. Verify that normal gameplay routes remain unaffected and that direct unauthorized navigation cannot retrieve dashboard data."""),
("9", "Build analytics query services and aggregate jobs", "Dashboard", "Fast, typed dashboard queries with consistent metric definitions.", """Implement the data-access layer for the Command Centre.

Create approved SQL views, materialized views, aggregate tables, or RPC functions for overview metrics, game performance, funnels, multiplayer health, and reliability. Every query must accept a date range and appropriate filters. Use the formulas in the measurement plan and return typed response shapes.

Add a safe aggregation or refresh mechanism suitable for Vercel and Supabase. Make refresh operations idempotent and prevent overlapping jobs. Show data freshness timestamps. Common dashboard requests should avoid scanning the full raw event table.

Add indexes based on actual query patterns, explain query plans for the heaviest endpoints, and enforce staff-only access. Create seeded test fixtures covering zero data, partial funnels, duplicate events, cross-day sessions, and development versus production data. Add reconciliation checks that compare aggregates with raw source events."""),
("10", "Build the Overview dashboard", "Dashboard", "A 60-second view of Decked's product health.", """Build the Command Centre Overview page using the approved query services and existing dashboard primitives.

Include daily, weekly, and monthly active players; new and returning players; game starts; completions; completion rate; cards played; median session duration; multiplayer rooms created and started; multiplayer start conversion; replay or rematch rate; and error-free session rate.

Provide date range, comparison period, game, play mode, device class, and environment filters where supported. Show previous-period deltas with accessible wording and avoid implying significance from very small samples. Include trend charts, top-games ranking, data freshness, metric-definition tooltips, loading states, zero-data states, partial-data warnings, and error states.

Preserve filters in the URL. Add tests for metric formatting, filter serialization, loading, empty results, and query failures. Validate desktop and tablet layouts and run an accessibility check."""),
("11", "Build Games and game detail analytics", "Dashboard", "Comparable performance data across all 21 games.", """Build the Games section of the Command Centre.

Create a sortable and filterable table for all 21 canonical games showing game views, selections, pass-and-play starts, multiplayer starts, unique players, first-card rate, completion rate, median cards played, median duration, replay or rematch rate, and seven-day or 30-day trend.

Selecting a game must open a detail page with trends, mode split, device split, funnel, deck or category performance, session depth, replay behavior, and relevant errors. Use the canonical registry so names and identifiers cannot drift from the product.

Support date filters, comparison periods, URL persistence, pagination if required, accessible table semantics, and CSV export of the currently filtered data. Do not expose prompt text or raw player identifiers. Add tests for sorting, filtering, export consistency, and games with no data."""),
("12", "Build funnel analytics", "Dashboard", "Clear measurement of discovery and multiplayer drop-off.", """Build the Funnels section using the canonical event definitions.

Implement at least three funnels:
1. Landing to browse to game selected to play mode selected to game started to first card to completion.
2. Play Together selected to room created to first guest joined to game started to first card to completion.
3. Join opened to code submitted to room joined to game started to first card to completion.

Show entrants, conversion from the previous step, conversion from the first step, and median elapsed time between steps. Support filters for game, play mode, device class, country when available, acquisition source, app version, environment, and date range.

Define how sessions are ordered, how repeat events are handled, and what conversion window is used. Make those rules visible in the interface. Add tests for out-of-order events, repeated steps, missing steps, cross-day sessions, and timezone boundaries."""),
("13", "Build Multiplayer and Reliability views", "Dashboard", "Operational visibility into room success, failures, and release quality.", """Build the Multiplayer and Reliability sections.

Multiplayer must show rooms created, successful joins, join failure reasons, lobby-to-start conversion, median time to first guest, median time to start, room-size distribution, completion rate, host disconnects, host-handoff success, Realtime failures, rematch requests, and rematches started. Add a privacy-safe room timeline for authorized staff if the technical design includes it.

Reliability must show frontend errors, failed RPC calls, Realtime status failures, error-free sessions, affected sessions, and trends by game, browser, device, app version, and release marker. Group identical errors without hiding occurrence counts or first and last seen times.

Do not reveal display names, room codes, free-form content, access tokens, stack-frame secrets, or raw anonymous IDs. Add loading, empty, redacted, and partial-data states. Test authorization, redaction, grouping, filtering, and failure conditions."""),
("14", "Add CSV exports and audit logging", "Quality and launch", "Useful exports that preserve privacy and accountability.", """Implement CSV export for Overview tables, Games, Funnels, Multiplayer, and Reliability where appropriate.

Exports must reflect the active filters and date range, use stable column names, include the generation timestamp and data-freshness timestamp, escape CSV safely, and remain usable for large result sets. Generate sensitive exports through a trusted server path and enforce staff authorization.

Record staff audit-log entries for exports, role changes, settings changes, and any room-level investigation. Do not log exported data or secrets in the audit entry. Build a simple admin-only audit viewer with date, actor, action, target type, and metadata summary.

Add tests for unauthorized exports, formula injection protection, filtered-data consistency, large exports, and audit-log creation."""),
("15", "Complete privacy and security hardening", "Quality and launch", "Documented and tested protections for player and staff data.", """Perform a focused privacy and security hardening pass on the complete Command Centre implementation.

Review Supabase RLS, grants, RPC security-definer functions, search_path handling, staff role assignment, service-role usage, browser bundles, environment variables, event properties, exports, logs, and error payloads. Attempt horizontal and vertical privilege escalation using test accounts.

Add automated tests proving that anonymous players cannot read raw analytics, aggregates, staff users, or audit logs; viewers cannot modify roles or settings; and browser clients cannot assign themselves staff privileges. Confirm no service-role key is bundled into client code.

Create a data inventory and retention document. Add a deletion or anonymization procedure for a known analytics identity. Verify that display names, room codes, free-form answers, prompt text, and credentials are rejected or redacted. Fix all high- and critical-severity findings and report remaining risks by severity."""),
("16", "Verify analytics accuracy end to end", "Quality and launch", "Evidence that dashboard numbers match real user journeys.", """Create and run an end-to-end analytics verification plan for Decked.

Exercise representative journeys in a non-production environment: browsing and starting a pass-and-play game, abandoning before the first card, completing a local game, creating and joining a multiplayer room, completing a multiplayer game, failed join, host disconnect and handoff, room end, replay, and rematch.

For every journey, compare the expected event sequence with stored raw events, derived sessions, aggregates, and dashboard values. Verify timestamps, session boundaries, game IDs, play modes, deduplication, environment separation, and privacy redaction. Test at least three different games, including one with specialized shared state.

Create an automated reconciliation script or test suite where practical. Produce a concise verification report with pass or fail status, discrepancies, fixes made, and any known measurement limitations. Do not approve launch while unexplained metric discrepancies exceed 1 percent in controlled fixtures."""),
("17", "Optimize performance and operational resilience", "Quality and launch", "A dashboard and event pipeline that do not degrade gameplay.", """Profile and harden the analytics pipeline and Command Centre before launch.

Measure analytics-client impact on page load and gameplay interactions, ingestion failure behavior, event batch size, database write throughput, aggregate refresh duration, dashboard query latency, and the largest export. Confirm that analytics requests never block primary gameplay actions.

Review query plans and add or adjust indexes only when supported by evidence. Introduce pagination, caching, materialization, or background refresh where needed. Add reasonable rate limits and payload-size limits to ingestion. Test retry storms, duplicate submissions, temporary Supabase failure, slow queries, and stale aggregates.

Set explicit service targets for event delivery, dashboard freshness, query response time, and error rate. Document capacity assumptions and operational recovery steps. Run a final production build and report measured results before and after optimization."""),
("18", "Prepare deployment launch and rollback", "Quality and launch", "A controlled production release with monitoring and recovery steps.", """Prepare the Decked Command Centre for production deployment without deploying until explicitly authorized.

Create a release checklist covering environment variables, Supabase migration order, RLS verification, staff bootstrap, aggregate scheduling, Vercel configuration, analytics enablement, production versus development separation, smoke tests, monitoring, retention jobs, and backups.

Create a staged rollout plan: database foundation, disabled client code, staff-only dashboard, internal event validation, limited analytics enablement, and full enablement. Define go or no-go checks at each stage. Provide a rollback plan that preserves captured analytics while disabling ingestion or dashboard access safely.

Add release markers so dashboard metrics can be compared across application versions. Define post-launch checks for the first hour, first day, and first week. Produce a launch report template containing metric coverage, data freshness, ingestion errors, dashboard errors, performance, privacy incidents, and owner actions.

Run the complete lint, type-check, test, build, migration-validation, security, and end-to-end verification suite. Report commands and results. Stop before production mutation and clearly state the exact approval needed to proceed."""),
]

for num, title_text, phase, outcome, body in prompts:
    doc.add_page_break()
    doc.add_heading(f"Prompt {num}  {title_text}", level=1)
    p=doc.add_paragraph()
    r=p.add_run("Phase  "); r.bold=True
    p.add_run(phase)
    p=doc.add_paragraph()
    r=p.add_run("Outcome  "); r.bold=True
    p.add_run(outcome)
    doc.add_heading("Copy and paste prompt", level=2)
    doc.add_paragraph(body, style="Prompt Text")
    doc.add_heading("Completion gate", level=2)
    gate = {
        "Discovery": "Do not proceed until the output matches the actual repository and all unresolved product decisions are listed explicitly.",
        "Data foundation": "Do not proceed until the change is tested, privacy constraints are verified, and existing gameplay still builds and works.",
        "Dashboard": "Do not proceed until authorization, loading, empty, error, and filtered states are verified with realistic fixtures.",
        "Quality and launch": "Do not treat the task as complete without evidence from tests, reconciliation, security review, and documented rollback steps.",
    }[phase]
    doc.add_paragraph(gate)

doc.add_page_break()
doc.add_heading("Optional follow up prompts", level=1)
optional = [
    ("Content management", "Extend the Decked Command Centre with versioned game, deck, category, and prompt management. Design a draft, review, publish, and rollback workflow. Preserve stable content IDs so historical analytics remains valid. Require previews, validation, audit logs, role-based publishing, and safe cache invalidation. Do not overwrite content directly in application source files without first presenting the migration strategy."),
    ("Automated alerts", "Add configurable operational alerts for sudden drops in game starts or completion, elevated join failures, Realtime failures, ingestion delays, and aggregate refresh failures. Use minimum sample thresholds and comparison windows to reduce noise. Keep notification destinations and credentials server-side, record alert history, and provide test and mute controls."),
    ("Experimentation", "Design and implement a privacy-safe experiment framework for Decked. Support deterministic anonymous assignment, exposure events, mutually exclusive variants, start and end dates, guardrail metrics, and result interpretation. Prevent experiments from changing game rules in ways that split multiplayer participants within the same room."),
    ("Scheduled reporting", "Add a weekly internal product report generated from the same canonical metric queries as the Command Centre. Include product health, top movements, multiplayer health, reliability, data-quality warnings, and release annotations. Do not create an external notification or schedule until the destination, recipients, and authorization are explicitly approved."),
]
for title_text, body in optional:
    doc.add_heading(title_text, level=2)
    doc.add_paragraph(body, style="Prompt Text")

doc.add_page_break()
doc.add_heading("Final launch checklist", level=1)
for text in [
    "Metric definitions are approved and versioned.",
    "All 21 games appear in the instrumentation coverage matrix.",
    "Development and production data are separated.",
    "Anonymous players cannot access analytics or staff data.",
    "No service-role credential appears in browser code or build output.",
    "Display names, room codes, prompt text, and free-form answers are absent from analytics.",
    "Controlled end-to-end journeys reconcile within the agreed tolerance.",
    "Dashboard loading, empty, partial, error, and stale-data states are verified.",
    "Migrations, rollout, monitoring, and rollback steps are documented.",
    "Production deployment has explicit human approval.",
]:
    p=doc.add_paragraph(style="List Bullet")
    p.add_run(text)

doc.add_paragraph()
p=doc.add_paragraph()
r=p.add_run("Practical rule  "); r.bold=True
p.add_run("If a prompt produces a large change that cannot be reviewed comfortably, stop and ask the coding agent to split it into smaller commits before continuing.")

doc.save(OUT)
print(OUT)
