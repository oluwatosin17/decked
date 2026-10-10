from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = 'Decked_Command_Centre_Metrics_Charts_PRD_and_Prompt_Pack.docx'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.68)
sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5); styles['Normal'].font.color.rgb = RGBColor(34,34,38)
styles['Title'].font.name = 'Aptos Display'; styles['Title'].font.size = Pt(31); styles['Title'].font.bold = True; styles['Title'].font.color.rgb = RGBColor(0,0,0)
for name, size in [('Heading 1',21),('Heading 2',15),('Heading 3',12)]:
    styles[name].font.name = 'Aptos Display'; styles[name].font.size = Pt(size); styles[name].font.bold = True; styles[name].font.color.rgb = RGBColor(0,0,0)
    styles[name].paragraph_format.space_before = Pt(13); styles[name].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.12

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(cell):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.first_child_found_in('w:tcBorders')
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); b.append(e)

def table(headers, rows, widths=None, font=8.5):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'25272C'); borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(font)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; c.text=str(val); borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(c,'F4F6F8')
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
                for r in p.runs: r.font.size=Pt(font)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def bullets(items):
    for item in items: doc.add_paragraph(item, style='List Bullet')

def prompt(num, title, text):
    doc.add_heading(f'Prompt {num} {title}', level=2)
    p=doc.add_paragraph()
    r=p.add_run(text.strip()); r.font.name='Aptos'; r.font.size=Pt(9.5)

# Cover
p=doc.add_paragraph('DECKED  INTERNAL PRODUCT DOCUMENT', style='Subtitle'); p.runs[0].font.color.rgb=RGBColor(202,47,44); p.runs[0].font.bold=True
doc.add_paragraph('Decked Command Centre Metrics and Charts Product Requirements', style='Title')
doc.add_paragraph('Product requirements and implementation prompt pack', style='Subtitle')
doc.add_paragraph('Version 1.0  |  9 October 2026  |  Owner Product and Analytics')
doc.add_paragraph('\nThis document defines what the Decked Command Centre must show, how every metric is calculated, which chart answers each product question, and the order in which implementation prompts should be given to a coding agent. The recommended MVP focuses on trustworthy product health, game performance, funnels, multiplayer operations, and reliability. Raw player content and direct identifiers remain outside the dashboard.')
doc.add_paragraph('\nDecision summary', style='Heading 2')
bullets(['Use Supabase authoritative facts for multiplayer and typed browser events for discovery and pass and play.', 'Build charts only from staff-only aggregate query services; never chart raw events in the browser.', 'Preserve date, comparison, game, play mode, device, and environment filters in the URL.', 'Treat freshness, partial coverage, small samples, empty data, and query failures as first-class UI states.', 'Do not launch until controlled fixtures reconcile within 1 percent and privacy tests pass.'])
doc.add_page_break()

doc.add_heading('Product definition', level=1)
doc.add_heading('Problem', level=2)
doc.add_paragraph('Decked needs one internal view that answers whether people discover games, begin playing, finish sessions, replay, succeed in multiplayer, and encounter product failures. The current product generates browser and database facts, but staff need consistent definitions and visual explanations rather than isolated event counts.')
doc.add_heading('Primary users', level=2)
table(['User','Decisions supported'], [('Founder and product lead','Overall health, adoption, retention, launch and prioritization'),('Game designer','Which games, modes and decks create depth, completion and replay'),('Engineering and operations','Join failures, Realtime issues, RPC errors, stale data and regressions'),('Growth and content','Discovery conversion, acquisition quality and device or release differences')],[1.7,5.2])
doc.add_heading('Goals', level=2)
bullets(['Explain product health in 60 seconds.', 'Compare all 21 canonical games without identifier drift.', 'Reveal discovery and multiplayer drop-off.', 'Separate production, preview and development data.', 'Make every number traceable to a documented formula and freshness timestamp.'])
doc.add_heading('Non goals', level=2)
bullets(['Player surveillance or room-code lookup.', 'Displaying names, answers, prompts, access tokens or raw anonymous IDs.', 'Real-time second-by-second monitoring in the MVP.', 'Statistical significance claims without sample and experiment design.'])

doc.add_heading('Information architecture', level=1)
table(['Section','Primary question','Default view'], [('Overview','Is Decked healthy and growing?','KPI strip, active-player trend, starts and completions, top games'),('Games','Which games perform best and why?','Comparable table and game detail'),('Funnels','Where do users stop?','Three ordered funnels with elapsed time'),('Multiplayer','Do rooms successfully become completed games?','Room conversion, timing, failures and rematches'),('Reliability','Which failures affect sessions and releases?','Error-free sessions, grouped failures and trends'),('Settings and audit','Who accessed or exported data?','Freshness, definitions and staff audit log')],[1.35,3.25,2.3])

doc.add_page_break()
doc.add_heading('Global controls and behavior', level=1)
table(['Control','Requirement'], [('Date range','Preset 7, 30 and 90 days plus custom inclusive UTC dates'),('Comparison','Previous period by default; same period previous year optional later'),('Environment','Production default; explicit preview and development selection'),('Game','All games or one canonical game ID'),('Play mode','All, pass and play, or play together'),('Device','All, mobile, tablet, desktop where observed'),('URL persistence','Every supported filter serialized and restorable from the URL'),('Freshness','Show last successful aggregate refresh and warn after 30 minutes'),('Small samples','Avoid directional language below the metric minimum sample threshold')],[1.6,5.3])
doc.add_heading('Metric presentation rules', level=2)
bullets(['Show current value, previous-period value, absolute or percentage delta, sample size when useful, and a plain-language definition.', 'Use zero only for a measured zero. Use unavailable for missing coverage and partial for incomplete dimensions.', 'Rates use summed numerators and denominators across the range, never averages of daily rates.', 'All day boundaries are UTC in storage; the interface states the timezone.', 'Charts must have text summaries, accessible legends, visible focus, non-color encodings, and tabular alternatives for exportable data.'])

doc.add_page_break()
doc.add_heading('Overview metrics and charts', level=1)
table(['Metric','Definition','Recommended visual'], [
('Daily active players','Distinct analytics user IDs with a qualifying event per UTC day','Line chart'),('Weekly and monthly active players','Distinct IDs in trailing 7 and 30 day windows','Three-line active-player chart'),('New and returning players','First qualifying event in range versus activity before range','Stacked columns'),('Game starts','Distinct sessions with a local or authoritative multiplayer start','KPI plus line'),('Completions','Distinct sessions reaching completion','KPI plus line'),('Completion rate','Completions divided by starts','KPI plus line'),('Cards played','Sum of cards presented in started sessions','KPI and area chart'),('Median session duration','Median completed_at minus started_at','KPI and line'),('Rooms created and started','Durable rooms created and rooms reaching first start','Two-series line'),('Multiplayer start conversion','Started rooms divided by rooms with a successful guest','KPI and line'),('Replay or rematch rate','Replay selections plus rematches divided by completions','KPI and line'),('Error-free session rate','Sessions without frontend, RPC or failed Realtime events','KPI and line'),('Top games','Rank by starts with completion and trend context','Horizontal bar ranking')
],[1.6,3.35,1.95],8.2)

doc.add_page_break()
doc.add_heading('Games metrics and charts', level=1)
table(['Table column','Meaning'], [('Game views','Play-mode or game-detail views attributed to the game'),('Selections','Committed selection from browse, Quick Play or recommendation'),('Pass and play starts','Local sessions that reach the first playable state'),('Multiplayer starts','Authoritative multiplayer sessions started'),('Unique players','Distinct pseudonymous browser installations where observable'),('First-card rate','Sessions showing a first card divided by setups'),('Completion rate','Completed sessions divided by started sessions'),('Median cards played','Median cards presented per started session'),('Median duration','Median duration for completed sessions'),('Replay or rematch rate','Linked replay or rematch actions divided by completions'),('Trend','Start change over 7 or 30 days')],[1.9,5.0])
doc.add_heading('Game detail charts', level=2)
table(['Chart','Design','Question answered'], [('Starts and completions','Two-line daily trend','Is interest translating into completed play?'),('Mode split','100 percent stacked bar plus counts','Is usage local or multiplayer?'),('Device split','Horizontal bars','Does performance vary by device?'),('Game funnel','Ordered bars with prior and first-step conversion','Where does this game lose players?'),('Deck or category performance','Sortable bars or table','Which stable modes drive depth?'),('Session depth','Histogram buckets for cards or rounds','How deeply do sessions progress?'),('Replay behavior','Completion to replay/rematch flow','Does the game create another round?'),('Relevant errors','Grouped table with trend sparkline','Which failures affect this game?')],[1.55,2.45,2.9],8.3)

doc.add_heading('Funnel analytics', level=1)
doc.add_paragraph('Every funnel shows entrants, conversion from the previous step, conversion from the first step, median elapsed time between steps, the conversion window, ordering rule and repeated-event rule.')
table(['Funnel','Steps'], [('Discovery','Landing → Browse → Game selected → Play mode selected → Game started → First card → Completion'),('Multiplayer host','Play Together selected → Room created → First guest joined → Game started → First card → Completion'),('Multiplayer guest','Join opened → Code submitted → Room joined → Game started → First card → Completion')],[1.7,5.2])
doc.add_paragraph('Recommended visual: an accessible vertical step chart on desktop and a table on narrow screens. Use one consistent bar scale, show counts directly, and never imply users are unique when the funnel is session-based.')

doc.add_page_break()
doc.add_heading('Multiplayer metrics and charts', level=1)
table(['Metric group','Metrics','Visual'], [('Volume','Rooms created, joins, starts, completions','Daily multi-line trend'),('Conversion','Lobby-to-start and multiplayer completion','KPI plus conversion bars'),('Timing','Median time to first guest and start','KPI plus percentile trend'),('Room size','Peak player-count distribution','Histogram'),('Join failures','Normalized reason and occurrence count','Horizontal bars and table'),('Host continuity','Disconnects, eligible handoffs, successful handoffs','KPI and trend'),('Realtime','Timed out, channel error and closed states','Stacked columns'),('Rematch','Requests and rematches started','KPI and conversion line')],[1.5,3.6,1.7],8.2)
doc.add_paragraph('Room-level timelines remain out of MVP unless separately approved. Aggregate views must not return room codes, display names, participant IDs or privacy-safe room references that enable staff linkage.')

doc.add_heading('Reliability metrics and charts', level=1)
table(['Metric','Definition or grouping','Visual'], [('Error-free sessions','Sessions without recorded failure events','KPI and daily line'),('Affected sessions','Distinct sessions with at least one failure','KPI and line'),('Frontend errors','Fingerprint and normalized class','Grouped table'),('Failed RPC calls','RPC name, error code and failure class','Stacked bars'),('Realtime failures','Channel type and failure status','Stacked bars'),('Release quality','Error-free rate by app version or release marker','Version comparison'),('First and last seen','Bounds for each grouped failure','Table columns')],[1.55,3.55,1.7],8.3)
doc.add_heading('Data states', level=2)
table(['State','Required response'], [('Loading','Skeletons preserve final layout; filters remain understandable'),('Zero data','Explain that no rows match the filters; offer reset'),('Partial','Name the missing dimension or source and avoid fabricated values'),('Stale','Show last refresh time and warning after 30 minutes'),('Error','Keep filters, provide retry and show no misleading previous result'),('Unauthorized','Return no dashboard payload and present staff access guidance')],[1.25,5.65])

doc.add_page_break()
doc.add_heading('Data contract and acceptance criteria', level=1)
bullets(['All dashboard reads require active staff authorization enforced by database functions or a trusted server path.', 'Common dashboard requests read aggregates rather than scanning the complete raw event table.', 'Every response includes refreshed_at and uses canonical game IDs.', 'Development and production data cannot mix in a query or comparison.', 'Duplicate event IDs and idempotency keys contribute once.', 'Exports exactly match active filters and neutralize spreadsheet formulas.', 'Display names, room codes, answers, prompt text, tokens and raw anonymous IDs never appear in responses or logs.', 'Controlled fixtures reconcile raw events, derived sessions, aggregates and displayed values within 1 percent.', 'Desktop and tablet layouts meet WCAG AA contrast and keyboard navigation requirements.'])
doc.add_heading('MVP release scope', level=2)
doc.add_paragraph('Release Overview, Games, Funnels, Multiplayer and Reliability with date, comparison, game, play-mode and environment filters. Include URL persistence, CSV export, freshness, definitions, staff-only access and all non-happy-path states. Defer custom dashboards, alert destinations, experiments, cohort builders and room-level investigations.')
doc.add_heading('Success criteria', level=2)
table(['Outcome','Target'], [('Trust','No unexplained controlled discrepancy above 1 percent'),('Freshness','Normal aggregate age at or below 15 minutes; warning after 30'),('Query performance','Dashboard p95 below 1 second for supported ranges'),('Reliability','Analytics-induced gameplay error rate below 0.1 percent'),('Privacy','Zero forbidden-content findings in automated payload and response tests'),('Usability','A product lead can identify the leading issue or opportunity within 60 seconds')],[1.8,5.1])

doc.add_page_break()
doc.add_heading('Implementation prompt pack', level=1)
doc.add_paragraph('Give these prompts in order. Each prompt assumes the previous completion gate has passed. Keep the repository instructions attached to every implementation request.')
base = 'Work in the existing Decked repository. Inspect current code, migrations, tests and repository instructions first. Preserve unrelated changes. State current behavior and intended change, identify security, privacy, migration and compatibility risks, implement the smallest complete change, add tests, run relevant lint, type-check, tests and build, and report files, verification, assumptions and rollback. '
prompt(1,'Validate metric contracts', base + '''Audit the implemented Command Centre metrics against docs/command-centre/measurement-plan-v1.md and the current SQL RPCs. Produce a metric dictionary for Overview, Games, Funnels, Multiplayer and Reliability with numerator, denominator, cohort, source of truth, freshness, supported filters, partial-coverage rules and minimum sample guidance. Do not change UI yet. List every mismatch between the PRD and repository and recommend the smallest correction. Completion gate: no chart work until definitions and discrepancies are explicit.''')
prompt(2,'Create the chart design system', base + '''Create reusable, accessible chart primitives using the project’s existing dependencies and visual language. Before adding a chart library, verify whether one is already installed and compare bundle cost, accessibility and maintenance. Support line, bar, stacked bar, histogram, funnel steps, legends, tooltips, tabular fallback, loading, empty, partial, stale and error states. Use dark neutral surfaces and restrained red accents. Add keyboard and screen-reader tests. Do not connect production queries yet.''')
prompt(3,'Build Overview charts', base + '''Implement Overview KPI and chart presentation from the approved aggregate RPCs. Include active-player trends, new versus returning players, starts and completions, cards, duration, room conversion, replay/rematch, error-free sessions and top games. Preserve filters in the URL. Show comparison deltas, sample-size caution, metric definitions and freshness. Add fixture-based tests for formatting, loading, zero, stale, partial and error states.''')
prompt(4,'Build Games comparison charts', base + '''Implement the all-games analytics table for the 21 canonical games and connect sortable metrics, trend display and CSV export. Use the canonical game registry. Add the game-detail route with starts/completions, mode split, device split, funnel, category or deck performance, session depth, replay behavior and relevant errors. Test sorting, filters, URL restoration, missing data and export consistency.''')
prompt(5,'Build funnel visualizations', base + '''Implement the three approved funnels using canonical event definitions: discovery, multiplayer host and multiplayer guest. Display entrants, prior-step conversion, first-step conversion and median elapsed time. Make ordering, repeat handling and conversion windows visible. Test out-of-order events, repeated steps, cross-day sessions, timezone boundaries and partial funnels.''')
prompt(6,'Build Multiplayer charts', base + '''Implement Multiplayer charts for room volume, successful joins, start conversion, timing, room-size distribution, completion, disconnects, host handoffs, Realtime failures and rematches. Use only aggregate or redacted responses. Do not expose room codes, display names, participant IDs or room references. Add authorization, redaction, stale-data and failure tests.''')
prompt(7,'Build Reliability charts', base + '''Implement Reliability charts and grouped failure tables for frontend errors, RPC failures, Realtime failures, error-free sessions, affected sessions, browser, device, app version and release marker. Preserve occurrence counts plus first and last seen times. Add grouping, filtering, redaction and high-cardinality safeguards.''')
prompt(8,'Add cross-dashboard interaction quality', base + '''Standardize filters, comparison language, tooltips, legends, focus behavior, responsive layouts and URL serialization across every Command Centre page. Add accessible table alternatives and ensure charts do not rely on color alone. Validate desktop and tablet layouts and run automated accessibility checks plus manual keyboard verification.''')
prompt(9,'Verify data accuracy and performance', base + '''Run controlled non-production journeys across at least three games and reconcile expected events, raw rows, sessions, aggregates, RPC responses and rendered dashboard values. Profile dashboard queries, aggregate refresh, chart rendering and largest export. Capture query plans before adding indexes. Do not approve launch if unexplained discrepancies exceed 1 percent or p95 dashboard queries exceed the approved service target.''')
doc.add_page_break()
prompt(10,'Prepare launch and monitoring', base + '''Create the Command Centre launch checklist covering migrations, staff access, environment variables, analytics enablement, aggregate scheduling, backups, rollback, privacy review, freshness monitoring and post-launch checks for the first hour, day and week. Add release markers and a launch-report template. Keep launch blocked until security, reconciliation, performance and accessibility evidence is attached.''')

doc.add_heading('Product decisions still required', level=1)
bullets(['Minimum sample threshold for deltas and rankings.', 'Whether weekly and monthly active players are trailing windows or calendar periods in labels.', 'Whether viewers may export every aggregate report.', 'Whether country segmentation is approved and available server-side.', 'Retention limit for aggregates and rate-limit counters.', 'Whether room-level investigations are allowed for administrators.', 'Final dashboard freshness schedule and operational owner.'])

footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run('Decked Command Centre Metrics and Charts Product Requirements'); run.font.size=Pt(8); run.font.color.rgb=RGBColor(100,100,105)
doc.core_properties.title='Decked Command Centre Metrics and Charts Product Requirements'
doc.core_properties.subject='Metrics, charts and implementation prompts for the Decked Command Centre'
doc.core_properties.author='Decked'
doc.save(OUT)
print(OUT)
